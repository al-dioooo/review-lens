from __future__ import annotations

import os
import shutil
import tempfile
from dataclasses import asdict, dataclass
from pathlib import Path
from uuid import uuid4

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from reviewlens.exceptions import ExportError
from reviewlens.export.charts import ChartArtifact, generate_charts
from reviewlens.export.safety import safe_excel_value, sanitize_frame
from reviewlens.models import AnalysisResult, Diagnostic, ExportManifest

_SHEET_ORDER = [
    "metadata",
    "reviews",
    "clusters",
    "summary",
    "keywords",
    "visual_data",
    "visuals",
]
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="4C78A8")
_TITLE_FILL = PatternFill(fill_type="solid", fgColor="274060")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_BODY_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
_HEADER_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
_MAX_COLUMN_WIDTH = 50
_IMAGE_WIDTH = 720
_APPROXIMATE_ROW_PIXELS = 20


@dataclass(frozen=True, slots=True)
class _VisualSection:
    title_row: int
    header_row: int
    data_end_row: int
    max_column: int


def _metadata_frame(result: AnalysisResult) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "key": list(result.metadata),
            "value": list(result.metadata.values()),
        }
    )


def _sanitize_frames(
    result: AnalysisResult,
) -> tuple[pd.DataFrame, dict[str, pd.DataFrame], dict[str, pd.DataFrame]]:
    tables = {
        "reviews": result.reviews,
        "clusters": result.clusters,
        "summary": result.summary,
        "keywords": result.keywords,
    }
    sanitized_tables: dict[str, pd.DataFrame] = {}
    sanitized_visual_data: dict[str, pd.DataFrame] = {}
    truncation_count = 0

    metadata, count = sanitize_frame(_metadata_frame(result))
    truncation_count += count
    for name, frame in tables.items():
        sanitized_tables[name], count = sanitize_frame(frame)
        truncation_count += count
    for name, frame in result.visual_data.items():
        sanitized_visual_data[name], count = sanitize_frame(frame)
        truncation_count += count

    if truncation_count:
        diagnostic = Diagnostic(
            severity="warning",
            code="export_truncation",
            message=("Excel export truncated cell values to 32,767 characters."),
            count=truncation_count,
        )
        value, _ = safe_excel_value(asdict(diagnostic))
        metadata = pd.concat(
            [
                metadata,
                pd.DataFrame(
                    {
                        "key": ["diagnostic.export_truncation"],
                        "value": [value],
                    }
                ),
            ],
            ignore_index=True,
        )
    return metadata, sanitized_tables, sanitized_visual_data


def _style_header(sheet: Worksheet, row: int, max_column: int) -> None:
    for cell in sheet.iter_cols(
        min_col=1,
        max_col=max_column,
        min_row=row,
        max_row=row,
    ):
        cell[0].fill = _HEADER_FILL
        cell[0].font = _HEADER_FONT
        cell[0].alignment = _HEADER_ALIGNMENT
    sheet.row_dimensions[row].height = 24


def _style_title(sheet: Worksheet, row: int, max_column: int) -> None:
    for cell in sheet.iter_cols(
        min_col=1,
        max_col=max_column,
        min_row=row,
        max_row=row,
    ):
        cell[0].fill = _TITLE_FILL
        cell[0].font = _HEADER_FONT
        cell[0].alignment = _HEADER_ALIGNMENT
    sheet.row_dimensions[row].height = 24


def _style_body(sheet: Worksheet, *, start_row: int = 2) -> None:
    for row in sheet.iter_rows(min_row=start_row):
        for cell in row:
            cell.alignment = _BODY_ALIGNMENT


def _fit_columns(sheet: Worksheet) -> None:
    for column_index in range(1, sheet.max_column + 1):
        maximum = 0
        for cell in sheet.iter_cols(
            min_col=column_index,
            max_col=column_index,
            min_row=1,
            max_row=sheet.max_row,
        ):
            for item in cell:
                if item.value is not None:
                    maximum = max(maximum, len(str(item.value)))
        width = min(max(maximum + 2, 10), _MAX_COLUMN_WIDTH)
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def _style_table_sheet(sheet: Worksheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    _style_body(sheet)
    _style_header(sheet, 1, sheet.max_column)
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    _fit_columns(sheet)


def _write_visual_data(
    writer: pd.ExcelWriter,
    frames: dict[str, pd.DataFrame],
) -> list[_VisualSection]:
    start_row = 0
    sections: list[_VisualSection] = []
    for name, frame in frames.items():
        frame.to_excel(
            writer,
            sheet_name="visual_data",
            startrow=start_row + 1,
            index=False,
        )
        sheet = writer.sheets["visual_data"]
        title_row = start_row + 1
        header_row = start_row + 2
        data_end_row = header_row + len(frame.index)
        sheet.cell(row=title_row, column=1, value=name)
        sections.append(
            _VisualSection(
                title_row=title_row,
                header_row=header_row,
                data_end_row=data_end_row,
                max_column=max(len(frame.columns), 1),
            )
        )
        start_row += len(frame.index) + 4
    return sections


def _style_visual_data(
    sheet: Worksheet,
    sections: list[_VisualSection],
) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    _style_body(sheet, start_row=1)
    for section in sections:
        _style_title(sheet, section.title_row, section.max_column)
        _style_header(sheet, section.header_row, section.max_column)
    if sections:
        first = sections[0]
        sheet.auto_filter.ref = (
            f"A{first.header_row}:"
            f"{get_column_letter(first.max_column)}{first.data_end_row}"
        )
    _fit_columns(sheet)


def _write_visuals(
    sheet: Worksheet,
    artifacts: tuple[ChartArtifact, ...],
) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    row = 1
    for artifact in artifacts:
        sheet.cell(row=row, column=1, value=artifact.title)
        _style_title(sheet, row, 1)
        image = OpenpyxlImage(str(artifact.path))
        if image.width > _IMAGE_WIDTH:
            ratio = _IMAGE_WIDTH / image.width
            image.width = _IMAGE_WIDTH
            image.height = round(image.height * ratio)
        image.anchor = f"A{row + 1}"
        sheet.add_image(image)
        row += max(round(image.height / _APPROXIMATE_ROW_PIXELS), 1) + 3
    sheet.column_dimensions["A"].width = _MAX_COLUMN_WIDTH


def _write_workbook(
    path: Path,
    metadata: pd.DataFrame,
    tables: dict[str, pd.DataFrame],
    visual_data: dict[str, pd.DataFrame],
    artifacts: tuple[ChartArtifact, ...],
) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        metadata.to_excel(writer, sheet_name="metadata", index=False)
        for name in ("reviews", "clusters", "summary", "keywords"):
            tables[name].to_excel(writer, sheet_name=name, index=False)
        sections = _write_visual_data(writer, visual_data)
        visuals = writer.book.create_sheet("visuals")

        for name in ("metadata", "reviews", "clusters", "summary", "keywords"):
            _style_table_sheet(writer.sheets[name])
        _style_visual_data(writer.sheets["visual_data"], sections)
        _write_visuals(visuals, artifacts)


def _validate_workbook(
    path: Path,
    artifacts: tuple[ChartArtifact, ...],
) -> None:
    for artifact in artifacts:
        if not artifact.path.is_file() or artifact.path.stat().st_size == 0:
            raise OSError(f"staged chart is missing or empty: {artifact.path.name}")
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if workbook.sheetnames != _SHEET_ORDER:
            raise OSError("staged workbook has an invalid sheet order")
        if len(workbook["visuals"]._images) != len(artifacts):
            raise OSError("staged workbook has an invalid embedded-image count")
    finally:
        workbook.close()


def _path_exists(path: Path) -> bool:
    return path.exists() or path.is_symlink()


def _remove_path(path: Path) -> None:
    if path.is_symlink() or path.is_file():
        path.unlink(missing_ok=True)
    elif path.is_dir():
        shutil.rmtree(path)


def _replace_path(source: Path, destination: Path) -> None:
    if _path_exists(destination):
        raise FileExistsError(destination)
    if source.is_dir() and not source.is_symlink():
        source.replace(destination)
    else:
        os.replace(source, destination)


def _backup_path(target: Path, token: str) -> Path:
    return target.with_name(f".{target.name}.{token}.reviewlens-backup")


def _install_staged_targets(
    staged_workbook: Path,
    staged_chart_dir: Path,
    output: Path,
    final_chart_dir: Path,
    *,
    force: bool,
) -> None:
    token = uuid4().hex
    workbook_backup = _backup_path(output, token)
    chart_backup = _backup_path(final_chart_dir, token)
    backed_up_workbook = False
    backed_up_charts = False
    installed_workbook = False
    installed_charts = False
    try:
        if not force and (_path_exists(output) or _path_exists(final_chart_dir)):
            raise FileExistsError("an export target already exists")
        if force and _path_exists(output):
            _replace_path(output, workbook_backup)
            backed_up_workbook = True
        if force and _path_exists(final_chart_dir):
            _replace_path(final_chart_dir, chart_backup)
            backed_up_charts = True

        if _path_exists(output) or _path_exists(final_chart_dir):
            raise FileExistsError("an export target already exists")
        os.replace(staged_workbook, output)
        installed_workbook = True
        staged_chart_dir.replace(final_chart_dir)
        installed_charts = True
    except Exception:
        backed_up_workbook = backed_up_workbook or _path_exists(workbook_backup)
        backed_up_charts = backed_up_charts or _path_exists(chart_backup)
        if backed_up_workbook:
            _remove_path(output)
            _replace_path(workbook_backup, output)
        elif installed_workbook:
            _remove_path(output)
        if backed_up_charts:
            _remove_path(final_chart_dir)
            _replace_path(chart_backup, final_chart_dir)
        elif installed_charts:
            _remove_path(final_chart_dir)
        raise
    else:
        _remove_path(workbook_backup)
        _remove_path(chart_backup)


def export_analysis(
    result: AnalysisResult,
    path: str | Path | None = None,
    *,
    force: bool = False,
) -> ExportManifest:
    """Export ``result`` to an atomically installed Excel report and PNGs."""
    requested = (
        Path.cwd() / f"{result.source_stem}-reviewlens.xlsx"
        if path is None
        else Path(path)
    )
    output = requested.resolve()
    if output.suffix.lower() != ".xlsx":
        raise ExportError("Excel export requires a .xlsx output path.")
    final_chart_dir = output.parent / f"{output.stem}_charts"
    existing = [target for target in (output, final_chart_dir) if _path_exists(target)]
    if existing and not force:
        names = ", ".join(target.name for target in existing)
        raise ExportError(f"Cannot export {output.name}; already exists: {names}.")

    staging: Path | None = None
    try:
        staging = Path(
            tempfile.mkdtemp(
                prefix=f".{output.stem}-reviewlens-",
                dir=output.parent,
            )
        )
        staged_workbook = staging / output.name
        staged_chart_dir = staging / final_chart_dir.name
        artifacts = generate_charts(result, staged_chart_dir)
        metadata, tables, visual_data = _sanitize_frames(result)
        _write_workbook(
            staged_workbook,
            metadata,
            tables,
            visual_data,
            artifacts,
        )
        _validate_workbook(staged_workbook, artifacts)
        _install_staged_targets(
            staged_workbook,
            staged_chart_dir,
            output,
            final_chart_dir,
            force=force,
        )
    except ExportError:
        raise
    except Exception as error:
        raise ExportError(f"Failed to export {output.name}: {error}") from error
    finally:
        if staging is not None:
            shutil.rmtree(staging, ignore_errors=True)

    chart_paths = tuple(final_chart_dir / artifact.path.name for artifact in artifacts)
    return ExportManifest(workbook_path=output, chart_paths=chart_paths)


__all__ = ["export_analysis"]
