from __future__ import annotations

import ast
import hashlib
import inspect
import re
import shlex
import shutil
import subprocess
import sys
from dataclasses import fields
from pathlib import Path
from typing import cast

import pytest
from openpyxl import load_workbook

import reviewlens
from reviewlens import AnalysisConfig, AnalysisResult, analyze_reviews
from reviewlens.config import (
    ClusteringConfig,
    IngestionConfig,
    InterpretationConfig,
    PreprocessingConfig,
    ProjectionConfig,
    VectorizerConfig,
)
from reviewlens.exceptions import AnalysisError
from reviewlens.ingestion.normalization import normalize_column_name

ROOT = Path(__file__).resolve().parents[1]
REQUIRED = (
    "README.md",
    "LICENSE",
    "NOTICE",
    "CONTRIBUTING.md",
    "CODE_OF_CONDUCT.md",
    "SECURITY.md",
    "CHANGELOG.md",
    "docs/api.md",
    "docs/architecture.md",
    "docs/data-format.md",
)
_LOCAL_LINK = re.compile(r"\]\((?![a-z][a-z0-9+.-]*:)([^)#]+)(?:#[^)]*)?\)")
_FENCE = re.compile(r"```(?P<language>[a-z0-9_-]+)\n(?P<body>.*?)\n```", re.DOTALL)
_README_QUICK_CLI = "reviewlens analyze examples/sample_reviews.csv"
_README_ADVANCED_CLI = (
    "reviewlens analyze reviews.csv --text-column text --rating-column rating "
    "--language id --clusters auto --output report.xlsx"
)
_README_PYTHON_API = """from reviewlens import analyze_reviews

result = analyze_reviews(
    file="reviews.csv",
    language="id",
    clusters=5
)

result.export_excel(
    "report.xlsx"
)"""


def _fenced_blocks(document: Path, language: str) -> tuple[str, ...]:
    content = document.read_text(encoding="utf-8")
    return tuple(
        match.group("body")
        for match in _FENCE.finditer(content)
        if match.group("language") == language
    )


def _readme_block(language: str, expected: str) -> str:
    matches = tuple(
        block
        for block in _fenced_blocks(ROOT / "README.md", language)
        if block == expected
    )
    assert len(matches) == 1
    return matches[0]


def _markdown_section(document: Path, heading: str) -> str:
    content = document.read_text(encoding="utf-8")
    start = content.index(f"## {heading}")
    following = content.find("\n## ", start + len(heading) + 3)
    return content[start:] if following == -1 else content[start:following]


def _documented_sheets() -> tuple[str, ...]:
    output = _markdown_section(ROOT / "README.md", "Output")
    return tuple(re.findall(r"^\d+\. `([^`]+)`", output, flags=re.MULTILINE))


def _run_documented_cli(command: str, cwd: Path) -> subprocess.CompletedProcess[str]:
    arguments = shlex.split(command)
    assert arguments[0] == "reviewlens"
    suffix = ".exe" if sys.platform == "win32" else ""
    executable = Path(sys.executable).with_name(f"reviewlens{suffix}")
    assert executable.is_file()
    return subprocess.run(
        [str(executable), *arguments[1:]],
        cwd=cwd,
        text=True,
        capture_output=True,
        check=False,
    )


def _assert_documented_export(workbook_path: Path, chart_directory: Path) -> None:
    assert workbook_path.is_file()
    workbook = load_workbook(workbook_path, read_only=False)
    try:
        assert tuple(workbook.sheetnames) == _documented_sheets()
    finally:
        workbook.close()
    assert chart_directory.is_dir()
    assert any(path.suffix == ".png" for path in chart_directory.iterdir())


def test_required_open_source_documents_exist_without_placeholders() -> None:
    forbidden = ("TBD", "TODO", "FIXME", "fill in")
    for relative in REQUIRED:
        content = (ROOT / relative).read_text(encoding="utf-8")
        assert content.strip(), relative
        assert not any(word.lower() in content.lower() for word in forbidden), relative


def test_readme_contains_required_user_sections() -> None:
    content = (ROOT / "README.md").read_text(encoding="utf-8")
    for heading in (
        "## Installation",
        "## Quick start",
        "## Python API",
        "## Output",
        "## Architecture",
        "## Roadmap",
        "## Contributing",
    ):
        assert heading in content


def test_license_and_security_identity() -> None:
    license_text = (ROOT / "LICENSE").read_text(encoding="utf-8")
    security = (ROOT / "SECURITY.md").read_text(encoding="utf-8")
    assert "Apache License" in license_text
    assert "Version 2.0, January 2004" in license_text
    assert "lisafronaldio123@gmail.com" in security


def test_canonical_policy_texts_are_unmodified_except_for_contact() -> None:
    expected_hashes = {
        "LICENSE": "cfc7749b96f63bd31c3c42b5c471bf756814053e847c10f3eb003417bc523d30",
        "CODE_OF_CONDUCT.md": (
            "99a08b9dd6f347f7610602fa1b79c3b55c7922cfea1c31569198d49f54e93a14"
        ),
    }
    for relative, expected in expected_hashes.items():
        digest = hashlib.sha256((ROOT / relative).read_bytes()).hexdigest()
        assert digest == expected, relative


def test_readme_examples_match_supported_contracts() -> None:
    assert _readme_block("bash", _README_QUICK_CLI) == _README_QUICK_CLI
    assert _readme_block("bash", _README_ADVANCED_CLI) == _README_ADVANCED_CLI
    assert _readme_block("python", _README_PYTHON_API) == _README_PYTHON_API


def test_readme_quick_cli_example_executes_exactly(tmp_path: Path) -> None:
    examples = tmp_path / "examples"
    examples.mkdir()
    shutil.copy2(ROOT / "examples/sample_reviews.csv", examples)
    command = _readme_block("bash", _README_QUICK_CLI)
    completed = _run_documented_cli(command, tmp_path)
    assert completed.returncode == 0, completed.stderr
    _assert_documented_export(
        tmp_path / "sample_reviews-reviewlens.xlsx",
        tmp_path / "sample_reviews-reviewlens_charts",
    )


def test_readme_advanced_cli_example_executes_exactly(tmp_path: Path) -> None:
    destination = tmp_path / "reviews.csv"
    shutil.copy2(ROOT / "examples/sample_reviews.csv", destination)
    source = destination.read_text(encoding="utf-8")
    assert source.startswith("reviewText,stars,")
    destination.write_text(
        source.replace("reviewText,stars,", "text,rating,", 1),
        encoding="utf-8",
    )
    command = _readme_block("bash", _README_ADVANCED_CLI)
    completed = _run_documented_cli(command, tmp_path)
    assert completed.returncode == 0, completed.stderr
    _assert_documented_export(
        tmp_path / "report.xlsx",
        tmp_path / "report_charts",
    )


def test_readme_python_api_example_executes_exactly(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    shutil.copy2(ROOT / "examples/sample_reviews.csv", tmp_path / "reviews.csv")
    monkeypatch.chdir(tmp_path)
    namespace: dict[str, object] = {}
    program = _readme_block("python", _README_PYTHON_API)
    exec(compile(program, "README.md", "exec"), namespace)
    result = cast(AnalysisResult, namespace["result"])
    assert result.metadata["selected_k"] == 5
    _assert_documented_export(
        tmp_path / "report.xlsx",
        tmp_path / "report_charts",
    )


def test_api_docs_track_public_exports_signature_and_config_sections() -> None:
    document = ROOT / "docs/api.md"
    content = document.read_text(encoding="utf-8")
    for exported in reviewlens.__all__:
        assert f"`{exported}`" in content

    signature_block = next(
        block
        for block in _fenced_blocks(document, "python")
        if block.startswith("def analyze_reviews(")
    )
    function = cast(ast.FunctionDef, ast.parse(signature_block).body[0])
    documented_parameters = [
        argument.arg
        for argument in (
            *function.args.posonlyargs,
            *function.args.args,
            *function.args.kwonlyargs,
        )
    ]
    assert documented_parameters == list(inspect.signature(analyze_reviews).parameters)

    for config_type in (
        IngestionConfig,
        PreprocessingConfig,
        VectorizerConfig,
        ClusteringConfig,
        InterpretationConfig,
        ProjectionConfig,
    ):
        assert f"`{config_type.__name__}`" in content
        for field in fields(config_type):
            assert f"`{field.name}=" in content


def test_api_docs_distinguish_prevalidated_and_downstream_config_errors(
    example_csv: Path,
) -> None:
    config = AnalysisConfig(vectorizer=VectorizerConfig(min_df=20))
    with pytest.raises(AnalysisError):
        analyze_reviews(example_csv, config=config)

    content = (ROOT / "docs/api.md").read_text(encoding="utf-8")
    assert "explicitly validated subset" in content
    assert re.search(r"`min_df`\s+or\s+`max_df`", content)
    assert "surface as `AnalysisError`" in content


def test_data_format_docs_match_ascii_only_column_normalization() -> None:
    assert normalize_column_name("Téks Ulasan") == "t_ks_ulasan"
    content = (ROOT / "docs/data-format.md").read_text(encoding="utf-8")
    assert "outside ASCII `[A-Za-z0-9]`" in content
    assert "non-ASCII letters are separators" in content


def test_link_matcher_includes_images_and_ignores_external_destinations() -> None:
    markdown = (
        "[document](docs/api.md)\n"
        "![local image](images/diagram.png)\n"
        "[![badge](https://img.example/badge.svg)](LICENSE)\n"
        "![remote image](https://img.example/remote.png)\n"
        "[email](mailto:maintainer@example.com)\n"
    )
    assert _LOCAL_LINK.findall(markdown) == [
        "docs/api.md",
        "images/diagram.png",
        "LICENSE",
    ]


def test_local_documentation_links_resolve() -> None:
    for relative in REQUIRED:
        if not relative.endswith(".md"):
            continue
        content = (ROOT / relative).read_text(encoding="utf-8")
        for target in _LOCAL_LINK.findall(content):
            resolved = (ROOT / relative).parent / target
            assert resolved.resolve().exists(), f"{relative}: {target}"
