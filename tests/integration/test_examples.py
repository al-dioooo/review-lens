import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from reviewlens import analyze_reviews


def test_seeded_example_is_reproducible(example_csv: Path) -> None:
    first = analyze_reviews(example_csv)
    second = analyze_reviews(example_csv)
    assert first.metadata["selected_k"] == second.metadata["selected_k"]
    assert first.reviews["cluster_id"].tolist() == second.reviews["cluster_id"].tolist()
    assert first.keywords.equals(second.keywords)
    assert first.clusters["representative_source_rows"].tolist() == (
        second.clusters["representative_source_rows"].tolist()
    )


def test_committed_csv_and_json_examples_are_equivalent(
    example_csv: Path,
    example_json: Path,
) -> None:
    csv_reviews = pd.read_csv(example_csv)
    payload = json.loads(example_json.read_text(encoding="utf-8"))
    json_reviews = payload["reviews"]
    fields = ["reviewText", "stars", "createdAt", "location"]

    assert csv_reviews.columns.tolist() == fields
    assert len(csv_reviews) == len(json_reviews) == 15
    assert pd.api.types.is_numeric_dtype(csv_reviews["stars"])
    assert all(
        isinstance(review["stars"], (int, float))
        and not isinstance(review["stars"], bool)
        for review in json_reviews
    )
    for field in fields:
        assert [review[field] for review in json_reviews] == (
            csv_reviews[field].tolist()
        )

    csv_result = analyze_reviews(example_csv, clusters=3)
    json_result = analyze_reviews(example_json, clusters=3)
    assert csv_result.reviews["review_text"].tolist() == (
        json_result.reviews["review_text"].tolist()
    )
    assert csv_result.reviews["cluster_id"].tolist() == (
        json_result.reviews["cluster_id"].tolist()
    )


def test_json_list_and_object_have_same_manual_analysis(
    tmp_path: Path,
    example_json: Path,
) -> None:
    payload = json.loads(example_json.read_text(encoding="utf-8"))
    list_path = tmp_path / "reviews-list.json"
    list_path.write_text(
        json.dumps(payload["reviews"], ensure_ascii=False),
        encoding="utf-8",
    )
    object_result = analyze_reviews(example_json, clusters=3)
    list_result = analyze_reviews(list_path, clusters=3)
    assert object_result.reviews["review_text"].tolist() == (
        list_result.reviews["review_text"].tolist()
    )
    assert object_result.reviews["cluster_id"].tolist() == (
        list_result.reviews["cluster_id"].tolist()
    )


def test_prd_report_destination_is_temporary(
    tmp_path: Path,
    example_csv: Path,
) -> None:
    result = analyze_reviews(example_csv, clusters=3)
    output = tmp_path / "examples" / "output" / "report.xlsx"
    output.parent.mkdir(parents=True)
    result.export_excel(output)
    workbook = load_workbook(output, read_only=False)
    assert workbook.sheetnames == [
        "metadata",
        "reviews",
        "clusters",
        "summary",
        "keywords",
        "visual_data",
        "visuals",
    ]
