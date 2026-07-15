from __future__ import annotations

import json
import os
from copy import deepcopy
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from reviewlens import AnalysisResult, Diagnostic, analyze_reviews
from reviewlens.exceptions import ExportError
from reviewlens.export.charts import ChartArtifact, generate_charts
from reviewlens.export.excel import _remove_path
from reviewlens.export.safety import sanitize_frame


def _source(tmp_path: Path, ratings: bool = True) -> Path:
    path = tmp_path / "reviews.csv"
    suffixes = [",2", ",1", ",2", ",5", ",4", ",5"] if ratings else [""] * 6
    texts = [
        "=pelayanan lambat antre",
        "antre lama pelayanan",
        "petugas lambat antre",
        "tempat bersih nyaman",
        "bersih rapi nyaman",
        "lokasi nyaman bersih",
    ]
    header = "text,rating\n" if ratings else "text\n"
    path.write_text(
        header
        + "".join(
            f"{text}{rating}\n" for text, rating in zip(texts, suffixes, strict=True)
        ),
        encoding="utf-8",
    )
    return path


def _existing_export(
    tmp_path: Path,
) -> tuple[AnalysisResult, Path, bytes, dict[str, bytes]]:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    output = tmp_path / "report.xlsx"
    manifest = result.export_excel(output)
    return (
        result,
        output,
        output.read_bytes(),
        {path.name: path.read_bytes() for path in manifest.chart_paths},
    )


def _chart_bytes(chart_dir: Path) -> dict[str, bytes]:
    return {path.name: path.read_bytes() for path in chart_dir.iterdir()}


def _assert_original_targets(
    output: Path,
    workbook: bytes,
    charts: dict[str, bytes],
) -> None:
    assert output.read_bytes() == workbook
    assert _chart_bytes(output.with_name(f"{output.stem}_charts")) == charts


def _assert_excel_rows(
    actual: list[tuple[object, ...]],
    expected: list[tuple[object, ...]],
) -> None:
    assert len(actual) == len(expected)
    for actual_row, expected_row in zip(actual, expected, strict=True):
        assert len(actual_row) == len(expected_row)
        for actual_value, expected_value in zip(
            actual_row,
            expected_row,
            strict=True,
        ):
            if isinstance(expected_value, float):
                assert actual_value == pytest.approx(expected_value)
            else:
                assert actual_value == expected_value


def test_export_has_ordered_sheets_embedded_images_and_pngs(tmp_path: Path) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    output = tmp_path / "report.xlsx"
    manifest = result.export_excel(output)
    assert manifest.workbook_path == output.resolve()
    assert manifest.workbook_path.is_file()
    assert manifest.chart_paths
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == [
        "metadata",
        "reviews",
        "clusters",
        "summary",
        "keywords",
        "visual_data",
        "visuals",
    ]
    headers = [cell.value for cell in workbook["reviews"][1]]
    review_column = headers.index("review_text") + 1
    assert workbook["reviews"].cell(2, review_column).value.startswith("'=")
    assert workbook["visuals"]._images


def test_default_path_uses_cwd_and_default_chart_directory(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    monkeypatch.chdir(tmp_path)
    manifest = result.export_excel()
    expected = tmp_path / "reviews-reviewlens.xlsx"
    assert manifest.workbook_path == expected
    assert manifest.workbook_path.is_file()
    assert manifest.chart_paths
    assert {path.parent for path in manifest.chart_paths} == {
        tmp_path / "reviews-reviewlens_charts"
    }


def test_chart_directory_alone_blocks_non_force_export(tmp_path: Path) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    output = tmp_path / "report.xlsx"
    chart_dir = tmp_path / "report_charts"
    chart_dir.mkdir()
    marker = chart_dir / "keep.txt"
    marker.write_text("original", encoding="utf-8")
    with pytest.raises(ExportError, match="already exists"):
        result.export_excel(output)
    assert not output.exists()
    assert marker.read_text(encoding="utf-8") == "original"


def test_workbook_tables_visual_blocks_and_styles(tmp_path: Path) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    output = tmp_path / "report.xlsx"
    result.export_excel(output)
    workbook = load_workbook(output, data_only=False)
    metadata = pd.DataFrame(
        {"key": list(result.metadata), "value": list(result.metadata.values())}
    )
    primary = {
        "metadata": metadata,
        "reviews": result.reviews,
        "clusters": result.clusters,
        "summary": result.summary,
        "keywords": result.keywords,
    }
    for name, source_frame in primary.items():
        expected, _ = sanitize_frame(source_frame)
        sheet = workbook[name]
        assert [cell.value for cell in sheet[1]] == expected.columns.tolist()
        _assert_excel_rows(
            list(sheet.iter_rows(min_row=2, values_only=True)),
            list(expected.itertuples(index=False, name=None)),
        )
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == (
            f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        )
        assert all(cell.font.bold for cell in sheet[1])
        assert all(cell.fill.fill_type == "solid" for cell in sheet[1])
        assert all(
            dimension.width is not None and dimension.width <= 50
            for dimension in sheet.column_dimensions.values()
        )
        if sheet.max_row > 1:
            assert sheet.cell(2, 1).alignment.wrap_text is True
            assert sheet.cell(2, 1).alignment.vertical == "top"

    visual_data = workbook["visual_data"]
    next_title_row = 1
    for name, source_frame in result.visual_data.items():
        expected, _ = sanitize_frame(source_frame)
        title_row = next_title_row
        header_row = title_row + 1
        assert visual_data.cell(title_row, 1).value == name
        assert visual_data.cell(title_row, 1).font.bold is True
        assert [
            visual_data.cell(header_row, column).value
            for column in range(1, len(expected.columns) + 1)
        ] == expected.columns.tolist()
        actual_rows = [
            tuple(
                visual_data.cell(row, column).value
                for column in range(1, len(expected.columns) + 1)
            )
            for row in range(header_row + 1, header_row + len(expected.index) + 1)
        ]
        _assert_excel_rows(
            actual_rows,
            list(expected.itertuples(index=False, name=None)),
        )
        first_blank = header_row + len(expected.index) + 1
        for blank_row in (first_blank, first_blank + 1):
            assert all(
                visual_data.cell(blank_row, column).value is None
                for column in range(1, max(len(expected.columns), 1) + 1)
            )
        next_title_row += len(expected.index) + 4
    assert visual_data.freeze_panes == "A2"
    assert visual_data.auto_filter.ref is not None

    visuals = workbook["visuals"]
    visual_titles = [cell.value for cell in visuals["A"] if cell.value is not None]
    assert len(visual_titles) == len(visuals._images)
    assert all(cell.font.bold for cell in visuals["A"] if cell.value is not None)
    assert visuals.freeze_panes == "A2"


def test_manifest_paths_and_visual_titles_preserve_chart_artifact_order(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    output = tmp_path / "report.xlsx"
    captured: tuple[ChartArtifact, ...] = ()

    def capture_artifacts(
        analysis: AnalysisResult,
        output_dir: Path,
    ) -> tuple[ChartArtifact, ...]:
        nonlocal captured
        captured = generate_charts(analysis, output_dir)
        return captured

    monkeypatch.setattr(
        "reviewlens.export.excel.generate_charts",
        capture_artifacts,
    )
    manifest = result.export_excel(output)
    assert [path.name for path in manifest.chart_paths] == [
        artifact.path.name for artifact in captured
    ]
    assert all(path.is_file() for path in manifest.chart_paths)
    workbook = load_workbook(output, data_only=False)
    titles = [cell.value for cell in workbook["visuals"]["A"] if cell.value is not None]
    assert titles == [artifact.title for artifact in captured]
    assert len(workbook["visuals"]._images) == len(captured)


def test_truncation_diagnostic_does_not_mutate_analysis_state(tmp_path: Path) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    long_review = "=" + ("x" * 40_000)
    result.reviews.loc[0, "review_text"] = long_review
    reviews_before = result.reviews.copy(deep=True)
    metadata_before = deepcopy(result.metadata)
    diagnostics_before = result.diagnostics
    output = tmp_path / "report.xlsx"
    result.export_excel(output)

    workbook = load_workbook(output, data_only=False)
    metadata_rows = {
        key: value
        for key, value in workbook["metadata"].iter_rows(
            min_row=2,
            values_only=True,
        )
    }
    diagnostic = json.loads(metadata_rows["diagnostic.export_truncation"])
    assert diagnostic == {
        "code": "export_truncation",
        "count": 1,
        "message": "Excel export truncated cell values to 32,767 characters.",
        "severity": "warning",
    }
    review_headers = [cell.value for cell in workbook["reviews"][1]]
    review_column = review_headers.index("review_text") + 1
    stored = workbook["reviews"].cell(2, review_column).value
    assert isinstance(stored, str)
    assert stored.startswith("'=")
    assert len(stored) == 32_767
    pd.testing.assert_frame_equal(result.reviews, reviews_before)
    assert result.metadata == metadata_before
    assert result.diagnostics is diagnostics_before


def test_export_does_not_mutate_any_analysis_result_component(tmp_path: Path) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    frame_names = ("reviews", "clusters", "summary", "keywords", "evaluation")
    frames_before = {
        name: getattr(result, name).copy(deep=True) for name in frame_names
    }
    visual_data_before = {
        name: frame.copy(deep=True) for name, frame in result.visual_data.items()
    }
    metadata_before = deepcopy(result.metadata)
    diagnostics_before: tuple[Diagnostic, ...] = result.diagnostics
    source_stem_before = result.source_stem
    result.export_excel(tmp_path / "report.xlsx")
    for name, expected in frames_before.items():
        pd.testing.assert_frame_equal(getattr(result, name), expected)
    assert set(result.visual_data) == set(visual_data_before)
    for name, expected in visual_data_before.items():
        pd.testing.assert_frame_equal(result.visual_data[name], expected)
    assert result.metadata == metadata_before
    assert result.diagnostics is diagnostics_before
    assert result.source_stem == source_stem_before


def test_export_refuses_existing_target_without_force(tmp_path: Path) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    output = tmp_path / "report.xlsx"
    result.export_excel(output)
    with pytest.raises(ExportError, match="already exists"):
        result.export_excel(output)
    result.export_excel(output, force=True)


def test_rating_free_export_omits_rating_chart(tmp_path: Path) -> None:
    result = analyze_reviews(_source(tmp_path, ratings=False), clusters=2)
    manifest = result.export_excel(tmp_path / "report.xlsx")
    assert not any(
        path.name == "rating_distribution.png" for path in manifest.chart_paths
    )


def test_force_export_restores_original_after_install_failure(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    output = tmp_path / "report.xlsx"
    first = result.export_excel(output)
    original_workbook = output.read_bytes()
    original_charts = {path.name: path.read_bytes() for path in first.chart_paths}
    real_replace = os.replace
    failed = False

    def fail_once(source: str | Path, destination: str | Path) -> None:
        nonlocal failed
        if Path(destination) == output and not failed:
            failed = True
            raise OSError("simulated staged-workbook install failure")
        real_replace(source, destination)

    monkeypatch.setattr("reviewlens.export.excel.os.replace", fail_once)
    with pytest.raises(ExportError, match=r"report\.xlsx"):
        result.export_excel(output, force=True)
    assert output.read_bytes() == original_workbook
    assert {
        path.name: path.read_bytes() for path in (tmp_path / "report_charts").iterdir()
    } == original_charts


def test_force_export_preserves_original_after_backup_failure(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    output = tmp_path / "report.xlsx"
    first = result.export_excel(output)
    original_workbook = output.read_bytes()
    original_charts = {path.name: path.read_bytes() for path in first.chart_paths}
    real_replace = os.replace

    def fail_workbook_backup(source: str | Path, destination: str | Path) -> None:
        if Path(source) == output and "reviewlens-backup" in Path(destination).name:
            raise OSError("simulated workbook backup failure")
        real_replace(source, destination)

    monkeypatch.setattr("reviewlens.export.excel.os.replace", fail_workbook_backup)
    with pytest.raises(ExportError, match=r"report\.xlsx"):
        result.export_excel(output, force=True)
    assert output.read_bytes() == original_workbook
    assert {
        path.name: path.read_bytes() for path in (tmp_path / "report_charts").iterdir()
    } == original_charts


def test_chart_backup_failure_restores_both_original_targets(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result, output, workbook, charts = _existing_export(tmp_path)
    chart_dir = tmp_path / "report_charts"
    real_replace = Path.replace
    failure = OSError("simulated chart backup failure")

    def fail_chart_backup(source: Path, destination: Path) -> Path:
        if source == chart_dir and "reviewlens-backup" in destination.name:
            raise failure
        return real_replace(source, destination)

    monkeypatch.setattr(Path, "replace", fail_chart_backup)
    with pytest.raises(ExportError) as caught:
        result.export_excel(output, force=True)
    assert caught.value.__cause__ is failure
    _assert_original_targets(output, workbook, charts)


def test_chart_install_failure_restores_both_original_targets(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result, output, workbook, charts = _existing_export(tmp_path)
    chart_dir = tmp_path / "report_charts"
    real_replace = Path.replace
    failure = OSError("simulated staged-chart install failure")

    def fail_chart_install(source: Path, destination: Path) -> Path:
        if source.name == chart_dir.name and destination == chart_dir:
            raise failure
        return real_replace(source, destination)

    monkeypatch.setattr(Path, "replace", fail_chart_install)
    with pytest.raises(ExportError) as caught:
        result.export_excel(output, force=True)
    assert caught.value.__cause__ is failure
    _assert_original_targets(output, workbook, charts)


def test_workbook_restore_failure_still_restores_charts_and_keeps_backup(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result, output, workbook, charts = _existing_export(tmp_path)
    chart_dir = tmp_path / "report_charts"
    real_path_replace = Path.replace
    real_os_replace = os.replace
    install_failure = OSError("simulated staged-chart install failure")

    def fail_chart_install(source: Path, destination: Path) -> Path:
        if source.name == chart_dir.name and destination == chart_dir:
            raise install_failure
        return real_path_replace(source, destination)

    def fail_workbook_restore(source: str | Path, destination: str | Path) -> None:
        if "reviewlens-backup" in Path(source).name and Path(destination) == output:
            raise OSError("simulated workbook restore failure")
        real_os_replace(source, destination)

    monkeypatch.setattr(Path, "replace", fail_chart_install)
    monkeypatch.setattr("reviewlens.export.excel.os.replace", fail_workbook_restore)
    with pytest.raises(ExportError) as caught:
        result.export_excel(output, force=True)
    assert caught.value.__cause__ is install_failure
    assert "rollback failures" in str(caught.value)
    assert "workbook restore" in str(caught.value)
    assert _chart_bytes(chart_dir) == charts
    workbook_backups = [
        path for path in tmp_path.glob(".*.reviewlens-backup") if path.is_file()
    ]
    assert len(workbook_backups) == 1
    assert workbook_backups[0].read_bytes() == workbook


def test_chart_restore_failure_still_restores_workbook_and_keeps_backup(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result, output, workbook, charts = _existing_export(tmp_path)
    chart_dir = tmp_path / "report_charts"
    real_replace = Path.replace
    install_failure = OSError("simulated staged-chart install failure")

    def fail_chart_install_and_restore(source: Path, destination: Path) -> Path:
        if source.name == chart_dir.name and destination == chart_dir:
            raise install_failure
        if "reviewlens-backup" in source.name and destination == chart_dir:
            raise OSError("simulated chart restore failure")
        return real_replace(source, destination)

    monkeypatch.setattr(Path, "replace", fail_chart_install_and_restore)
    with pytest.raises(ExportError) as caught:
        result.export_excel(output, force=True)
    assert caught.value.__cause__ is install_failure
    assert "rollback failures" in str(caught.value)
    assert "charts restore" in str(caught.value)
    assert output.read_bytes() == workbook
    chart_backups = [
        path for path in tmp_path.glob(".*.reviewlens-backup") if path.is_dir()
    ]
    assert len(chart_backups) == 1
    assert _chart_bytes(chart_backups[0]) == charts


def test_transient_backup_cleanup_failure_retries_both_and_returns_success(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result, output, _, _ = _existing_export(tmp_path)
    cleanup_calls: list[str] = []
    failed = False

    def fail_first_backup_cleanup(path: Path) -> None:
        nonlocal failed
        if "reviewlens-backup" in path.name:
            cleanup_calls.append(path.name)
            if not failed:
                failed = True
                raise OSError("simulated transient backup cleanup failure")
        _remove_path(path)

    monkeypatch.setattr(
        "reviewlens.export.excel._remove_path",
        fail_first_backup_cleanup,
    )
    manifest = result.export_excel(output, force=True)
    assert manifest.workbook_path == output
    assert len({name.split(".reviewlens-backup")[0] for name in cleanup_calls}) == 2
    assert not list(tmp_path.glob(".*.reviewlens-backup"))


def test_helper_export_error_is_wrapped_with_basename_and_cause(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    output = tmp_path / "report.xlsx"
    helper_error = ExportError("simulated helper validation error")

    def fail_validation(_path: Path, _artifacts: object) -> None:
        raise helper_error

    monkeypatch.setattr("reviewlens.export.excel._validate_workbook", fail_validation)
    with pytest.raises(ExportError) as caught:
        result.export_excel(output)
    assert caught.value is not helper_error
    assert str(caught.value).startswith("Failed to export report.xlsx:")
    assert caught.value.__cause__ is helper_error


def test_staged_validation_failure_installs_no_final_target(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    output = tmp_path / "report.xlsx"
    validation_error = OSError("simulated staged validation failure")

    def fail_validation(_path: Path, _artifacts: object) -> None:
        raise validation_error

    monkeypatch.setattr("reviewlens.export.excel._validate_workbook", fail_validation)
    with pytest.raises(ExportError) as caught:
        result.export_excel(output)
    assert caught.value.__cause__ is validation_error
    assert not output.exists()
    assert not (tmp_path / "report_charts").exists()


def test_export_requires_xlsx_extension(tmp_path: Path) -> None:
    result = analyze_reviews(_source(tmp_path), clusters=2)
    with pytest.raises(ExportError, match=r"\.xlsx"):
        result.export_excel(tmp_path / "report.xls")
