from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from reviewlens.cli import main


def _run(*arguments: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-m", "reviewlens", *arguments],
        text=True,
        capture_output=True,
        check=False,
    )


def _run_installed(*arguments: str) -> subprocess.CompletedProcess[str]:
    suffix = ".exe" if sys.platform == "win32" else ""
    executable = Path(sys.executable).with_name(f"reviewlens{suffix}")
    assert executable.is_file()
    return subprocess.run(
        [str(executable), *arguments],
        text=True,
        capture_output=True,
        check=False,
    )


def _source(tmp_path: Path) -> Path:
    path = tmp_path / "reviews.csv"
    path.write_text(
        "text,rating\n"
        "pelayanan lambat antre,2\n"
        "antre lama pelayanan,1\n"
        "petugas lambat antre,2\n"
        "tempat bersih nyaman,5\n"
        "bersih rapi nyaman,4\n"
        "lokasi nyaman bersih,5\n",
        encoding="utf-8",
    )
    return path


def _metadata_config(output: Path) -> dict[str, object]:
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        metadata = dict(
            workbook["metadata"].iter_rows(
                min_row=2,
                values_only=True,
            )
        )
    finally:
        workbook.close()
    value = metadata["config"]
    assert isinstance(value, str)
    parsed = json.loads(value)
    assert isinstance(parsed, dict)
    return parsed


def test_cli_analyze_writes_requested_report(tmp_path: Path) -> None:
    source = _source(tmp_path)
    output = tmp_path / "report.xlsx"
    completed = _run(
        "analyze",
        str(source),
        "--clusters",
        "2",
        "--output",
        str(output),
    )
    assert completed.returncode == 0
    assert output.is_file()
    assert "Analyzing:" in completed.stdout
    assert "Exporting workbook" in completed.stdout
    assert "Selected k: 2" in completed.stdout
    assert "Workbook:" in completed.stdout
    assert "Chart directory:" in completed.stdout
    assert completed.stderr == ""


def test_cli_input_error_is_exit_two_without_traceback(tmp_path: Path) -> None:
    completed = _run("analyze", str(tmp_path / "missing.csv"))
    assert completed.returncode == 2
    assert "error:" in completed.stderr.lower()
    assert "Traceback" not in completed.stderr


def test_cli_analysis_failure_is_exit_one(tmp_path: Path) -> None:
    source = tmp_path / "reviews.csv"
    source.write_text("text\nsama\nsama\nsama\n", encoding="utf-8")
    output = tmp_path / "report.xlsx"
    completed = _run("analyze", str(source), "--output", str(output))
    assert completed.returncode == 1
    assert "error:" in completed.stderr.lower()
    assert not output.exists()
    assert not (tmp_path / "report_charts").exists()


def test_cli_quiet_suppresses_success_output(tmp_path: Path) -> None:
    source = _source(tmp_path)
    output = tmp_path / "report.xlsx"
    completed = _run(
        "analyze",
        str(source),
        "--clusters",
        "auto",
        "--output",
        str(output),
        "--quiet",
    )
    assert completed.returncode == 0
    assert completed.stdout == ""
    assert completed.stderr == ""


def test_cli_quiet_duplicate_heavy_success_has_clean_stderr(tmp_path: Path) -> None:
    source = tmp_path / "duplicates.csv"
    source.write_text(
        "text,rating\n"
        "pelayanan lambat,1\n"
        "pelayanan lambat,1\n"
        "tempat bersih,5\n"
        "tempat bersih,5\n"
        "lokasi nyaman,4\n"
        "lokasi nyaman,4\n",
        encoding="utf-8",
    )
    output = tmp_path / "report.xlsx"

    completed = _run(
        "analyze",
        str(source),
        "--output",
        str(output),
        "--quiet",
    )

    assert completed.returncode == 0
    assert output.is_file()
    assert completed.stdout == ""
    assert completed.stderr == ""


@pytest.mark.parametrize("quiet", [False, True])
def test_cli_structured_review_text_is_exit_two_without_traceback(
    tmp_path: Path,
    quiet: bool,
) -> None:
    source = tmp_path / "structured.json"
    source.write_text(
        json.dumps(
            [
                {"text": ["nested", "review"], "rating": 5},
                {"text": "pelayanan cepat", "rating": 5},
                {"text": "pelayanan lambat", "rating": 1},
            ]
        ),
        encoding="utf-8",
    )
    arguments = ["analyze", str(source)]
    if quiet:
        arguments.append("--quiet")

    completed = _run(*arguments)

    assert completed.returncode == 2
    assert "error:" in completed.stderr.lower()
    assert "Traceback" not in completed.stderr
    if quiet:
        assert completed.stdout == ""


def test_cli_empty_csv_is_exit_two_without_traceback(tmp_path: Path) -> None:
    source = tmp_path / "empty.csv"
    source.write_bytes(b"")

    completed = _run("analyze", str(source), "--quiet")

    assert completed.returncode == 2
    assert completed.stdout == ""
    assert "error:" in completed.stderr.lower()
    assert "Traceback" not in completed.stderr


def test_installed_entrypoint_exposes_help_and_module_exposes_version() -> None:
    installed = _run_installed("--help")
    module = _run("--version")
    assert installed.returncode == 0
    assert "analyze" in installed.stdout
    assert installed.stderr == ""
    assert module.returncode == 0
    assert module.stdout.strip() == "reviewlens 0.1.0"
    assert module.stderr == ""


def test_cli_translates_all_analysis_and_export_flags(tmp_path: Path) -> None:
    source = tmp_path / "custom.csv"
    source.write_text(
        "body,score\n"
        "pelayanan lambat antre,2\n"
        "antre lama pelayanan,1\n"
        "petugas lambat antre,2\n"
        "tempat bersih nyaman,5\n"
        "bersih rapi nyaman,4\n"
        "lokasi nyaman bersih,5\n",
        encoding="utf-8",
    )
    output = tmp_path / "report.xlsx"
    output.write_bytes(b"old workbook")
    chart_directory = tmp_path / "report_charts"
    chart_directory.mkdir()
    (chart_directory / "old.txt").write_text("old charts", encoding="utf-8")

    completed = _run(
        "analyze",
        str(source),
        "--text-column",
        "body",
        "--rating-column",
        "score",
        "--language",
        "id",
        "--clusters",
        "2",
        "--output",
        str(output),
        "--ngram-range",
        "1",
        "1",
        "--min-df",
        "1",
        "--max-features",
        "25",
        "--stem",
        "--no-slang",
        "--no-stopwords",
        "--evaluate",
        "--force",
    )

    assert completed.returncode == 0
    assert completed.stderr == ""
    assert not (chart_directory / "old.txt").exists()
    assert list(chart_directory.glob("*.png"))
    config = _metadata_config(output)
    assert config["language"] == "id"
    preprocessing = config["preprocessing"]
    assert isinstance(preprocessing, dict)
    assert preprocessing["stem"] is True
    assert preprocessing["normalize_slang"] is False
    assert preprocessing["remove_stopwords"] is False
    vectorizer = config["vectorizer"]
    assert isinstance(vectorizer, dict)
    assert vectorizer["ngram_range"] == [1, 1]
    assert vectorizer["min_df"] == 1
    assert vectorizer["max_features"] == 25
    clustering = config["clustering"]
    assert isinstance(clustering, dict)
    assert clustering["clusters"] == 2
    assert clustering["evaluate_manual"] is True


def test_cli_configuration_error_is_exit_two(tmp_path: Path) -> None:
    source = _source(tmp_path)
    output = tmp_path / "report.xlsx"
    completed = _run(
        "analyze",
        str(source),
        "--ngram-range",
        "2",
        "1",
        "--output",
        str(output),
    )
    assert completed.returncode == 2
    assert "error:" in completed.stderr.lower()
    assert "Traceback" not in completed.stderr
    assert not output.exists()


def test_cli_argument_error_is_exit_two(tmp_path: Path) -> None:
    source = _source(tmp_path)
    completed = _run("analyze", str(source), "--clusters", "1")
    assert completed.returncode == 2
    assert "clusters must be at least 2" in completed.stderr
    assert "Traceback" not in completed.stderr


def test_cli_non_positive_vectorizer_argument_is_exit_two(tmp_path: Path) -> None:
    source = _source(tmp_path)
    completed = _run("analyze", str(source), "--min-df", "0")
    assert completed.returncode == 2
    assert "must be positive" in completed.stderr
    assert "Traceback" not in completed.stderr


def test_cli_export_failure_is_exit_one_without_partial_artifacts(
    tmp_path: Path,
) -> None:
    source = _source(tmp_path)
    output = tmp_path / "report.csv"
    completed = _run(
        "analyze",
        str(source),
        "--clusters",
        "2",
        "--output",
        str(output),
    )
    assert completed.returncode == 1
    assert "error:" in completed.stderr.lower()
    assert "Traceback" not in completed.stderr
    assert not output.exists()
    assert not (tmp_path / "report_charts").exists()


def test_cli_debug_includes_traceback_and_preserves_exit_code(
    tmp_path: Path,
) -> None:
    completed = _run(
        "analyze",
        str(tmp_path / "missing.csv"),
        "--debug",
    )
    assert completed.returncode == 2
    assert "Traceback (most recent call last):" in completed.stderr
    assert "InputError" in completed.stderr
    assert "error:" in completed.stderr.lower()


def test_cli_quiet_does_not_suppress_errors(tmp_path: Path) -> None:
    completed = _run(
        "analyze",
        str(tmp_path / "missing.csv"),
        "--quiet",
    )
    assert completed.returncode == 2
    assert completed.stdout == ""
    assert "error:" in completed.stderr.lower()


def test_main_returns_zero_after_in_process_analysis(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    source = _source(tmp_path)
    output = tmp_path / "direct.xlsx"
    exit_code = main(
        [
            "analyze",
            str(source),
            "--text-column",
            "text",
            "--rating-column",
            "rating",
            "--language",
            "id",
            "--clusters",
            "2",
            "--output",
            str(output),
            "--ngram-range",
            "1",
            "1",
            "--min-df",
            "1",
            "--max-features",
            "25",
            "--stem",
            "--no-slang",
            "--no-stopwords",
            "--evaluate",
            "--force",
        ]
    )
    captured = capsys.readouterr()
    assert exit_code == 0
    assert output.is_file()
    assert "Selected k: 2" in captured.out
    assert captured.err == ""


def test_main_returns_two_and_prints_debug_traceback(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    exit_code = main(
        [
            "analyze",
            str(tmp_path / "missing.csv"),
            "--debug",
        ]
    )
    captured = capsys.readouterr()
    assert exit_code == 2
    assert "Traceback (most recent call last):" in captured.err
    assert "error:" in captured.err.lower()


def test_main_returns_one_for_analysis_error(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    source = tmp_path / "degenerate.csv"
    source.write_text("text\nsama\nsama\nsama\n", encoding="utf-8")
    exit_code = main(["analyze", str(source), "--quiet"])
    captured = capsys.readouterr()
    assert exit_code == 1
    assert captured.out == ""
    assert "error:" in captured.err.lower()
